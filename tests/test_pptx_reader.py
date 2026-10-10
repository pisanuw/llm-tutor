"""Tests for PPTX ingestion."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

from pipeline.ingest.pptx_reader import (
    LectureData,
    extract_slide_content,
    save_lecture_json,
    _is_likely_code,
)


def _make_minimal_pptx(tmp_path: Path, title: str = "Test Lecture", body: str = "Some content") -> Path:
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    pptx_path = tmp_path / "test.pptx"
    prs.save(str(pptx_path))
    return pptx_path


def test_extract_slide_content_returns_lecture_data(tmp_path):
    pptx_path = _make_minimal_pptx(tmp_path)
    result = extract_slide_content(pptx_path)
    assert isinstance(result, LectureData)
    assert result.filename == "test.pptx"
    assert result.slide_count == 1
    assert result.slides[0].slide_number == 1


def test_extract_slide_title(tmp_path):
    pptx_path = _make_minimal_pptx(tmp_path, title="AST 101 Welcome")
    result = extract_slide_content(pptx_path)
    assert result.lecture_title == "AST 101 Welcome"
    assert result.slides[0].title == "AST 101 Welcome"


def test_save_lecture_json(tmp_path):
    pptx_path = _make_minimal_pptx(tmp_path)
    lecture = extract_slide_content(pptx_path)
    out = save_lecture_json(lecture, tmp_path / "out")
    assert out.exists()
    data = json.loads(out.read_text())
    assert data["filename"] == "test.pptx"
    assert data["slide_count"] == 1
    assert isinstance(data["slides"], list)


def test_full_text_includes_title_and_body(tmp_path):
    pptx_path = _make_minimal_pptx(tmp_path, title="Seasons", body="Axial tilt matters")
    lecture = extract_slide_content(pptx_path)
    assert "Seasons" in lecture.full_text
    assert "Axial tilt" in lecture.full_text


def test_is_likely_code():
    assert _is_likely_code("int main() { return 0; }") is True
    assert _is_likely_code("#include <iostream> void foo() {}") is True
    assert _is_likely_code("This is just regular text about the seasons.") is False


def test_empty_pptx(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])  # Blank slide
    pptx_path = tmp_path / "blank.pptx"
    prs.save(str(pptx_path))
    result = extract_slide_content(pptx_path)
    assert result.slide_count == 1
    assert result.slides[0].title == ""
