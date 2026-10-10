"""Tests for PPTX slide generation and writing."""

from __future__ import annotations


import pytest
from pptx import Presentation  # used for loading saved files

from pipeline.generate.pptx_writer import (
    parse_content_md,
    create_presentation,
    write_pptx,
)

SAMPLE_CONTENT_MD = """
---
## Slide 1: AST 101 Welcome
type: title
title: AST 101: The Sun, Moon & Stars
body: Week 1, Lecture 1
notes: Welcome students warmly.
---

---
## Slide 2: Agenda
type: agenda
title: Today's Agenda
bullets:
  - Course overview
  - Finding your way on the sky
  - Observing log 1
notes: Walk through each item briefly.
---

---
## Slide 3: Why Look Up?
type: content
title: Why Look Up?
bullets:
  - The sky as a clock and a calendar
  - Light as our only messenger
notes: Start with the sundial analogy.
---

---
## Slide 4: A Simple Example
type: code
title: Hello Flux
code: |
  int d = 2;
  double flux = 1.0 / (d * d);
  std::cout << flux << std::endl;
notes: Run it live and double the distance.
---

---
## Slide 5: Summary
type: summary
title: What We Covered
bullets:
  - Course expectations
  - First inverse-square program
notes: Preview next lecture.
---
"""


def test_parse_content_md_returns_slides():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    assert len(slides) == 5


def test_parse_slide_types():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    assert slides[0].slide_type == "title"
    assert slides[1].slide_type == "agenda"
    assert slides[2].slide_type == "content"
    assert slides[3].slide_type == "code"
    assert slides[4].slide_type == "summary"


def test_parse_slide_titles():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    assert slides[0].title == "AST 101: The Sun, Moon & Stars"
    assert slides[1].title == "Today's Agenda"


def test_parse_bullets():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    agenda = slides[1]
    assert "Course overview" in agenda.bullets
    assert "Finding your way on the sky" in agenda.bullets
    assert len(agenda.bullets) == 3


def test_parse_code_slide():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    code_slide = slides[3]
    assert "int d = 2;" in code_slide.code


def test_parse_notes():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    assert slides[0].notes == "Welcome students warmly."


def test_create_presentation_produces_pptx():
    slides = parse_content_md(SAMPLE_CONTENT_MD)
    prs = create_presentation(slides)
    # python-pptx Presentation is not a plain class; check duck-typing instead
    assert hasattr(prs, "slides")
    assert len(prs.slides) == 5


def test_write_pptx(tmp_path):
    lec_dir = tmp_path / "lecture-01"
    lec_dir.mkdir()
    (lec_dir / "content.md").write_text(SAMPLE_CONTENT_MD)
    out = write_pptx(1, tmp_path, verbose=False)
    assert out.exists()
    assert out.suffix == ".pptx"
    prs = Presentation(str(out))
    assert len(prs.slides) == 5


def test_write_pptx_missing_content(tmp_path):
    with pytest.raises(FileNotFoundError):
        write_pptx(99, tmp_path)


def test_parse_empty_string():
    slides = parse_content_md("")
    assert slides == []


def test_parse_no_slide_markers():
    slides = parse_content_md("Just some plain text without any slide headers.")
    assert slides == []
