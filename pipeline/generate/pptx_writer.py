"""Convert slide content.md to PowerPoint (.pptx) using python-pptx."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from tutor.courses import IDENTITY, get_course


@dataclass
class ParsedSlide:
    number: int
    slide_type: str
    title: str
    bullets: list[str] = field(default_factory=list)
    body: str = ""
    code: str = ""
    notes: str = ""


def parse_content_md(content_md: str) -> list[ParsedSlide]:
    slides: list[ParsedSlide] = []
    raw_slides = re.split(r"^---\s*$", content_md, flags=re.MULTILINE)

    slide_num = 0
    for block in raw_slides:
        block = block.strip()
        if not block:
            continue
        header_match = re.search(r"## Slide \d+:", block)
        if not header_match:
            continue
        slide_num += 1

        def _field(name: str) -> str:
            m = re.search(rf"^{name}:\s*(.+)$", block, re.MULTILINE)
            return m.group(1).strip() if m else ""

        def _multiline_field(name: str) -> str:
            m = re.search(rf"^{name}:\s*\|\s*\n(.*?)(?=^\w|\Z)", block, re.MULTILINE | re.DOTALL)
            return m.group(1).rstrip() if m else ""

        slide_type = _field("type") or "content"
        title = _field("title")

        bullets: list[str] = []
        bullets_match = re.search(r"^bullets:\s*\n(.*?)(?=^\w|\Z)", block, re.MULTILINE | re.DOTALL)
        if bullets_match:
            for line in bullets_match.group(1).splitlines():
                line = line.strip().lstrip("- ").strip()
                if line:
                    bullets.append(line)

        body = _multiline_field("body")
        code = _multiline_field("code")
        notes = _field("notes")

        slides.append(ParsedSlide(
            number=slide_num,
            slide_type=slide_type,
            title=title,
            bullets=bullets,
            body=body,
            code=code,
            notes=notes,
        ))

    return slides


def _new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    return prs


def title_subtitle(course: str = "") -> str:
    """The title slide's fallback subtitle: "AST 101 | Dr. Morgan", or the
    instructor alone when ``course`` is not a course in courses.yaml."""
    c = get_course(course) if course else None
    return f"{c.label} | {IDENTITY.instructor_short}" if c else IDENTITY.instructor_short


def _add_title_slide(prs: Presentation, slide: ParsedSlide, subtitle: str = "") -> None:
    layout = prs.slide_layouts[0]  # Title Slide
    sl = prs.slides.add_slide(layout)
    sl.shapes.title.text = slide.title
    if sl.placeholders and len(sl.placeholders) > 1:
        try:
            sl.placeholders[1].text = slide.body or subtitle or title_subtitle()
        except Exception:
            pass
    _set_notes(sl, slide.notes)


def _add_agenda_slide(prs: Presentation, slide: ParsedSlide) -> None:
    layout = prs.slide_layouts[1]  # Title and Content
    sl = prs.slides.add_slide(layout)
    sl.shapes.title.text = slide.title or "Today's Agenda"
    tf = sl.placeholders[1].text_frame
    tf.text = ""
    for i, bullet in enumerate(slide.bullets):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = bullet
        p.level = 0
    _set_notes(sl, slide.notes)


def _add_content_slide(prs: Presentation, slide: ParsedSlide) -> None:
    layout = prs.slide_layouts[1]
    sl = prs.slides.add_slide(layout)
    sl.shapes.title.text = slide.title

    content_ph = sl.placeholders[1]
    tf = content_ph.text_frame
    tf.word_wrap = True

    if slide.bullets:
        for i, bullet in enumerate(slide.bullets):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = bullet
            p.level = 0
            p.font.size = Pt(20)
    elif slide.body:
        tf.paragraphs[0].text = slide.body
        tf.paragraphs[0].font.size = Pt(18)

    _set_notes(sl, slide.notes)


def _add_code_slide(prs: Presentation, slide: ParsedSlide) -> None:
    layout = prs.slide_layouts[6]  # Blank
    sl = prs.slides.add_slide(layout)

    # Title text box at top
    title_box = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
    title_box.text_frame.paragraphs[0].text = slide.title
    title_box.text_frame.paragraphs[0].font.size = Pt(28)
    title_box.text_frame.paragraphs[0].font.bold = True

    # Code area with dark background
    code_box = sl.shapes.add_textbox(Inches(0.3), Inches(1.2), Inches(12.7), Inches(5.5))
    fill = code_box.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x1E, 0x1E, 0x1E)  # VS Code dark

    tf = code_box.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = slide.code or "// code here"
    run = p.runs[0] if p.runs else p.add_run()
    run.font.name = "Courier New"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0xD4, 0xD4, 0xD4)

    _set_notes(sl, slide.notes)


def _add_summary_slide(prs: Presentation, slide: ParsedSlide) -> None:
    layout = prs.slide_layouts[1]
    sl = prs.slides.add_slide(layout)
    sl.shapes.title.text = slide.title or "Summary"
    tf = sl.placeholders[1].text_frame
    tf.text = ""
    for i, bullet in enumerate(slide.bullets):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = bullet
        p.level = 0
    _set_notes(sl, slide.notes)


def _add_diagram_placeholder_slide(prs: Presentation, slide: ParsedSlide) -> None:
    layout = prs.slide_layouts[6]  # Blank
    sl = prs.slides.add_slide(layout)

    title_box = sl.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.8))
    title_box.text_frame.paragraphs[0].text = slide.title
    title_box.text_frame.paragraphs[0].font.size = Pt(28)
    title_box.text_frame.paragraphs[0].font.bold = True

    placeholder_box = sl.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    fill = placeholder_box.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF0)
    tf = placeholder_box.text_frame
    tf.text = f"[ DIAGRAM: {slide.body or slide.title} ]\n\n(Insert diagram here)"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    _set_notes(sl, slide.notes)


def _add_generic_slide(prs: Presentation, slide: ParsedSlide) -> None:
    _add_content_slide(prs, slide)


def _set_notes(pptx_slide, notes_text: str) -> None:
    if not notes_text:
        return
    try:
        notes_slide = pptx_slide.notes_slide
        notes_slide.notes_text_frame.text = notes_text
    except Exception:
        pass


SLIDE_WRITERS = {
    "title": _add_title_slide,
    "agenda": _add_agenda_slide,
    "content": _add_content_slide,
    "code": _add_code_slide,
    "summary": _add_summary_slide,
    "quiz": _add_summary_slide,
    "diagram-placeholder": _add_diagram_placeholder_slide,
}


def create_presentation(slides: list[ParsedSlide], subtitle: str = "") -> Presentation:
    """``subtitle`` fills a title slide that has no body of its own."""
    prs = _new_presentation()
    for slide in slides:
        if slide.slide_type == "title":
            _add_title_slide(prs, slide, subtitle)
            continue
        writer = SLIDE_WRITERS.get(slide.slide_type, _add_generic_slide)
        writer(prs, slide)
    return prs


def write_pptx(lecture_num: int, output_dir: Path, verbose: bool = False) -> Path:
    lec_dir = output_dir / f"lecture-{lecture_num:02d}"
    content_path = lec_dir / "content.md"
    if not content_path.exists():
        raise FileNotFoundError(f"content.md not found at {content_path}. Run 'slides' first.")

    content_md = content_path.read_text()
    slides = parse_content_md(content_md)
    if not slides:
        raise ValueError(f"No slides parsed from {content_path}")

    # output/<course>/lecture-NN: the folder above the lecture names the course.
    prs = create_presentation(slides, subtitle=title_subtitle(Path(output_dir).name))
    out_path = lec_dir / "slides.pptx"
    prs.save(str(out_path))

    if verbose:
        print(f"[pptx] Lecture {lecture_num:02d}: {len(slides)} slides -> {out_path}")
    return out_path
