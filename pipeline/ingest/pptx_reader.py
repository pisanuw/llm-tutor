"""Extract structured content from PowerPoint lecture files."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path

from pptx import Presentation


@dataclass
class SlideData:
    slide_number: int
    layout_name: str
    title: str
    body_text: list[str]
    speaker_notes: str
    has_code: bool = False


@dataclass
class LectureData:
    filename: str
    lecture_title: str
    slides: list[SlideData] = field(default_factory=list)

    @property
    def slide_count(self) -> int:
        return len(self.slides)

    @property
    def full_text(self) -> str:
        parts = []
        for s in self.slides:
            if s.title:
                parts.append(s.title)
            parts.extend(s.body_text)
            if s.speaker_notes:
                parts.append(s.speaker_notes)
        return "\n".join(parts)


def _shape_text(shape) -> str:
    try:
        if shape.has_text_frame:
            return shape.text_frame.text.strip()
    except Exception:
        pass
    return ""


def _is_likely_code(text: str) -> bool:
    code_markers = ["{", "}", "->", "<<", ">>", "::", "#include", "int ", "void ", "return ", "nullptr"]
    return sum(1 for m in code_markers if m in text) >= 2


def _extract_slide(slide, slide_number: int) -> SlideData:
    layout_name = ""
    try:
        layout_name = slide.slide_layout.name
    except Exception:
        pass

    title = ""
    body_texts: list[str] = []

    for shape in slide.shapes:
        text = _shape_text(shape)
        if not text:
            continue
        # Title placeholder is typically idx 0 or named "Title"
        is_title = False
        try:
            if shape.is_placeholder and shape.placeholder_format.idx == 0:
                is_title = True
        except Exception:
            pass
        if is_title and not title:
            title = text
        elif text != title:
            body_texts.append(text)

    notes = ""
    try:
        notes_slide = slide.notes_slide
        notes_tf = notes_slide.notes_text_frame
        notes = notes_tf.text.strip()
    except Exception:
        pass

    all_text = " ".join([title] + body_texts + [notes])
    has_code = _is_likely_code(all_text)

    return SlideData(
        slide_number=slide_number,
        layout_name=layout_name,
        title=title,
        body_text=body_texts,
        speaker_notes=notes,
        has_code=has_code,
    )


def extract_slide_content(pptx_path: Path) -> LectureData:
    prs = Presentation(str(pptx_path))
    slides_data: list[SlideData] = []
    for i, slide in enumerate(prs.slides, start=1):
        slides_data.append(_extract_slide(slide, i))

    lecture_title = pptx_path.stem
    if slides_data and slides_data[0].title:
        lecture_title = slides_data[0].title

    return LectureData(
        filename=pptx_path.name,
        lecture_title=lecture_title,
        slides=slides_data,
    )


def extract_all_lectures(lectures_dir: Path) -> dict[str, LectureData]:
    results: dict[str, LectureData] = {}
    pptx_files = sorted(lectures_dir.glob("*.pptx"))
    for pptx_path in pptx_files:
        results[pptx_path.name] = extract_slide_content(pptx_path)
    return results


def save_lecture_json(lecture: LectureData, output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(lecture.filename).stem
    out_path = output_dir / f"{stem}.json"
    payload = {
        "filename": lecture.filename,
        "lecture_title": lecture.lecture_title,
        "slide_count": lecture.slide_count,
        "slides": [asdict(s) for s in lecture.slides],
    }
    out_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False))
    return out_path


def ingest_lectures(lectures_dir: Path, output_dir: Path, verbose: bool = False) -> list[Path]:
    lectures = extract_all_lectures(lectures_dir)
    written: list[Path] = []
    for name, lecture in lectures.items():
        out = save_lecture_json(lecture, output_dir)
        written.append(out)
        if verbose:
            print(f"  {name}: {lecture.slide_count} slides -> {out.name}")
    return written
